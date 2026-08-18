"""Gemensamma delar för Excel-import: normalisering av telefonnummer och
personnummer, tolkning av celler samt rapporttyper.

Importen är tolerant: saknade kolumner rapporteras i stället för krasch, och
fel på enskilda rader samlas per rad medan giltiga rader importeras."""
from __future__ import annotations

import hashlib
import hmac
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.config import PNR_HASH_NYCKEL
from app.tid import till_utc


class Normaliseringsfel(ValueError):
    """Ett cellvärde kunde inte tolkas. Ger radfel i importrapporten."""


@dataclass
class Radfel:
    rad: int
    fel: str


@dataclass
class ImportRapport:
    saknade_kolumner: list[str] = field(default_factory=list)
    skapade: int = 0
    uppdaterade: int = 0
    hoppade_over: int = 0
    skapade_kunder: list[str] = field(default_factory=list)
    radfel: list[Radfel] = field(default_factory=list)
    # Sätts av atomära importer (bokningar): radfel → allt rullas tillbaka.
    aterrullad: bool = False

    @property
    def genomford(self) -> bool:
        """False om importen inte fullföljdes (saknade kolumner eller
        återrullad atomär import)."""
        return not self.saknade_kolumner and not self.aterrullad

    def sammanfattning(self) -> str:
        rader: list[str] = []
        if self.saknade_kolumner:
            rader.append(
                "Importen kunde inte köras — följande kolumner saknas i filen: "
                + ", ".join(self.saknade_kolumner)
            )
            return "\n".join(rader)
        if self.aterrullad:
            rader.append(
                f"VARNING: {len(self.radfel)} radfel — hela importen rullades "
                f"tillbaka, INGA rader importerades."
            )
            for fel in self.radfel:
                rader.append(f"  Rad {fel.rad}: {fel.fel}")
            rader.append(
                "ATL-kontrollen kan inte litas på förrän bokningsimporten är "
                "komplett. Rätta felen i filen och kör om importen."
            )
            return "\n".join(rader)
        rader.append(
            f"Import klar: {self.skapade} skapade, {self.uppdaterade} uppdaterade, "
            f"{self.hoppade_over} överhoppade dubbletter, {len(self.radfel)} radfel."
        )
        for namn in self.skapade_kunder:
            rader.append(f"  Ny kund skapades automatiskt: {namn}")
        for fel in self.radfel:
            rader.append(f"  Rad {fel.rad}: {fel.fel}")
        return "\n".join(rader)


def normalisera_telefon(varde: object) -> str:
    """Normalisera ett svenskt telefonnummer till E.164 (+46...).

    Hanterar 070..., +46 70..., 0046..., 46... samt numeriska Excel-celler
    där den inledande nollan tappats (t.ex. 701234567)."""
    if varde is None:
        raise Normaliseringsfel("telefonnummer saknas")
    if isinstance(varde, float) and varde.is_integer():
        s = str(int(varde))
    elif isinstance(varde, int):
        s = str(varde)
    else:
        s = str(varde).strip()
    s = re.sub(r"[\s\-().]", "", s)
    if s.startswith("00"):
        s = "+" + s[2:]
    if s.startswith("+46"):
        rest = s[3:]
    elif s.startswith("46") and len(s) >= 10:
        rest = s[2:]
    elif s.startswith("0"):
        rest = s[1:]
    elif s.startswith("7") and len(s) == 9:
        rest = s  # numerisk cell där Excel tappat inledande nolla
    else:
        raise Normaliseringsfel(f"ogiltigt telefonnummer: {varde!r}")
    if rest.startswith("0"):
        rest = rest[1:]
    if not rest.isdigit() or not (7 <= len(rest) <= 9):
        raise Normaliseringsfel(f"ogiltigt telefonnummer: {varde!r}")
    return "+46" + rest


def luhn_kontrollsiffra(niosiffror: str) -> int:
    """Kontrollsiffra enligt Luhn för personnummer (ÅÅMMDDNNN)."""
    summa = 0
    for i, tecken in enumerate(niosiffror):
        p = int(tecken) * (2 if i % 2 == 0 else 1)
        summa += p // 10 + p % 10
    return (10 - summa % 10) % 10


def _som_datum(ar: int, man: int, dag: int) -> date | None:
    if dag > 60:
        dag -= 60  # samordningsnummer
    try:
        return date(ar, man, dag)
    except ValueError:
        return None


def normalisera_personnummer(varde: object) -> str:
    """Normalisera personnummer till 12 siffror (ÅÅÅÅMMDDNNNN).

    Validerar datumdel (samordningsnummer med dag+60 accepteras) och
    Luhn-kontrollsiffra. Sekelregel för 10 siffror: 20xx om personen då är
    minst 16 år gammal är omöjligt väljs 19xx (konsulter antas vara ≥ 16 år).
    """
    if varde is None:
        raise Normaliseringsfel("personnummer saknas")
    s = re.sub(r"\D", "", str(varde).strip())
    if len(s) == 12:
        tolv = s
    elif len(s) == 10:
        idag = date.today()
        fodd_20 = _som_datum(2000 + int(s[:2]), int(s[2:4]), int(s[4:6]))
        minst_16 = fodd_20 is not None and fodd_20 <= date(idag.year - 16, idag.month, 1)
        tolv = ("20" if minst_16 else "19") + s
    else:
        raise Normaliseringsfel(f"personnummer ska ha 10 eller 12 siffror: {varde!r}")
    if _som_datum(int(tolv[:4]), int(tolv[4:6]), int(tolv[6:8])) is None:
        raise Normaliseringsfel(f"ogiltigt datum i personnummer: {varde!r}")
    if luhn_kontrollsiffra(tolv[2:11]) != int(tolv[11]):
        raise Normaliseringsfel(f"felaktig kontrollsiffra i personnummer: {varde!r}")
    return tolv


def hash_personnummer(personnummer12: str, nyckel: str | None = None) -> str:
    """HMAC-SHA256 av ett normaliserat personnummer. Klartext lagras aldrig."""
    n = nyckel if nyckel is not None else PNR_HASH_NYCKEL
    return hmac.new(n.encode(), personnummer12.encode(), hashlib.sha256).hexdigest()


def tolka_bool(varde: object, standard: bool = True) -> bool:
    if varde is None or (isinstance(varde, str) and not varde.strip()):
        return standard
    if isinstance(varde, bool):
        return varde
    s = str(varde).strip().lower()
    if s in ("ja", "j", "1", "true", "sant", "x"):
        return True
    if s in ("nej", "n", "0", "false", "falskt"):
        return False
    raise Normaliseringsfel(f"ogiltigt ja/nej-värde: {varde!r}")


def tolka_tid(varde: object) -> datetime:
    """Tolka en tidpunkt från en Excel-cell. Naiva tider antas vara
    Europe/Stockholm och konverteras till UTC."""
    if isinstance(varde, datetime):
        return till_utc(varde)
    if isinstance(varde, str):
        s = varde.strip()
        try:
            return till_utc(datetime.fromisoformat(s))
        except ValueError:
            pass
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H.%M", "%d/%m/%Y %H:%M"):
            try:
                return till_utc(datetime.strptime(s, fmt))
            except ValueError:
                continue
    raise Normaliseringsfel(f"ogiltig tidpunkt: {varde!r}")


def las_excel(sokvag: str | Path) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
    """Läs första arket. Returnerar (normaliserade kolumnnamn,
    [(radnummer, {kolumnnamn: värde})]). Tomma rader hoppas över."""
    wb = load_workbook(sokvag, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        raditer = ws.iter_rows(values_only=True)
        try:
            huvud = next(raditer)
        except StopIteration:
            return [], []
        kolumner = [str(c).strip().lower() if c is not None else "" for c in huvud]
        rader: list[tuple[int, dict[str, object]]] = []
        for radnr, rad in enumerate(raditer, start=2):
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in rad):
                continue
            rader.append(
                (
                    radnr,
                    {
                        kolumner[i]: rad[i]
                        for i in range(min(len(kolumner), len(rad)))
                        if kolumner[i]
                    },
                )
            )
        return kolumner, rader
    finally:
        wb.close()
